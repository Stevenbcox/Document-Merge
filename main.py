import os
import re
import fitz
import openpyxl
from PIL import Image
from reportlab.pdfgen import canvas
from datetime import datetime
import shelve
import pyodbc
from utils import progress_callback

# ----------------- CONFIG -----------------

ORDER = [
    "Terms",
    "Supplemental Borrower",
    "Bill Statement - CHARGE OFF",
    "Bill Statement",
    "Pay History",
    "Sales Memo",
    "GOODBYE",
    "OwnershipChain",
    "DataString",
]

SKIPPED_EXTENSIONS = ('.eml', '.htm', '.xlsx')
VALID_EXTENSIONS = ('.pdf', '.tif', '.tiff', '.jpg', '.jpeg')


# ----------------- SQL -----------------

def get_db_connection():
    """Connect to SQL Server using saved credentials."""
    with shelve.open(r'P:/Users/Justin/Projects/sql_creds/credentials') as db:
        server = db['server']
        database = db['database']
        username = db['username']
        password = db['password']

    conn_str = (
        "DRIVER=ODBC Driver 18 for SQL Server;"
        f"SERVER={server};DATABASE={database};"
        "ENCRYPT=no;"
        f"UID={username};PWD={password}"
    )
    return pyodbc.connect(conn_str)


def fetch_fileno_map(unique_numbers):
    """Map 9-digit numbers to 6-digit FILENO from SQL."""
    if not unique_numbers:
        return {}

    placeholders = ",".join("?" for _ in unique_numbers)
    query = f"""
        SELECT FORW_REFNO, FILENO
        FROM CLSMI.dbo.MASTER
        WHERE FORW_REFNO IN ({placeholders})
    """
    fileno_map = {}
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(query, list(unique_numbers))
        for forw_refno, fileno in cursor.fetchall():
            fileno_map[str(forw_refno)] = str(fileno)
    return fileno_map


# ----------------- UTILITIES -----------------

def get_unique_number(filename):
    """Extract 9-digit number from filename."""
    match = re.search(r'\d{9}', filename)
    return match.group() if match else None


def extract_date(filename):
    """Extract date from filename formatted as 'Stm. Date - mm_dd_yyyy'."""
    match = re.search(r'Stm\. Date - (\d{1,2}_\d{1,2}_\d{4})', filename)
    if match:
        return datetime.strptime(match.group(1), '%m_%d_%Y')
    return None


def get_sort_key(filename):
    """Determine sort priority based on ORDER list."""
    name = filename.lower()
    for index, keyword in enumerate(ORDER):
        if keyword.lower() in name:
            return index
    return len(ORDER)


def convert_image_to_pdf(image_path):
    """Convert JPG/JPEG to PDF."""
    img = Image.open(image_path)
    if img.width > img.height:
        img = img.rotate(90, expand=True)

    output_path = os.path.splitext(image_path)[0] + "_converted.pdf"
    c = canvas.Canvas(output_path, pagesize=img.size)
    c.drawImage(image_path, 0, 0, width=img.size[0], height=img.size[1])
    c.save()
    return output_path


def convert_tif_to_pdf(tif_path):
    """Convert TIFF/TIF to PDF."""
    img = Image.open(tif_path)
    if img.width > img.height:
        img = img.rotate(90, expand=True)

    output_path = os.path.splitext(tif_path)[0] + "_converted.pdf"
    img.save(output_path, "PDF", resolution=100.0)
    return output_path


def check_pdf_integrity(path):
    """Verify PDF can be opened."""
    try:
        with fitz.open(path) as doc:
            doc.load_page(0)
        return True
    except Exception:
        return False


def safe_rename(src, dest):
    """Rename file without overwriting."""
    if os.path.exists(dest):
        base, ext = os.path.splitext(dest)
        counter = 1
        while True:
            new_dest = f"{base}_{counter}{ext}"
            if not os.path.exists(new_dest):
                dest = new_dest
                break
            counter += 1
    os.rename(src, dest)


def write_results_excel(results, output_folder):
    """Write merge results to Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["rsg number", "merged"])
    for unique_number, merged in results:
        ws.append([unique_number, "x" if merged else ""])
    wb.save(os.path.join(output_folder, "merge_results.xlsx"))


# ----------------- FILE SCANNING -----------------

def scan_files(input_folder, filter_numbers=None):
    """Recursively scan folder and convert images/TIFFs to PDF."""
    pdf_dict = {}
    for root, _, files in os.walk(input_folder):
        for file in files:
            if file.lower().endswith(SKIPPED_EXTENSIONS):
                continue
            if not file.lower().endswith(VALID_EXTENSIONS):
                continue

            unique_number = get_unique_number(file)
            if not unique_number:
                continue
            if filter_numbers and unique_number not in filter_numbers:
                continue

            file_path = os.path.join(root, file)
            if file.lower().endswith(('.tif', '.tiff')):
                file_path = convert_tif_to_pdf(file_path)
            elif file.lower().endswith(('.jpg', '.jpeg')):
                file_path = convert_image_to_pdf(file_path)

            pdf_dict.setdefault(unique_number, []).append(
                (file_path, extract_date(file), file)
            )
    return pdf_dict


# ----------------- MERGING -----------------

def merge_documents(unique_number, documents, output_folder):
    """Merge PDFs for a single 9-digit number."""
    # Filter and sort
    documents = [d for d in documents if get_sort_key(d[2]) < len(ORDER)]
    documents.sort(key=lambda x: x[1] or datetime.min, reverse=True)
    documents.sort(key=lambda x: get_sort_key(x[2]))

    # Check integrity
    for path, _, _ in documents:
        if not check_pdf_integrity(path):
            return False, None

    output_path = os.path.join(output_folder, f"{unique_number}.pdf")

    with fitz.open() as merged_pdf:
        for path, _, _ in documents:
            with fitz.open(path) as src_pdf:
                for page in src_pdf:
                    if page.rotation == 0 and page.rect.width > page.rect.height:
                        page.set_rotation(90)
                merged_pdf.insert_pdf(src_pdf)
        if merged_pdf.page_count == 0:
            return False, None
        merged_pdf.save(output_path)

    return True, output_path


# ----------------- MAIN -----------------

def main(input_folder, output_folder, unique_numbers_list, progress_queue=None):
    """Scan input folder, filter by user 9-digit numbers, merge, rename, and save results."""
    # Scan recursively and filter
    pdf_dict = scan_files(input_folder, set(unique_numbers_list))
    if not pdf_dict:
        print("No matching documents found.")
        return

    total_documents = sum(len(docs) for docs in pdf_dict.values()) or 1
    processed = 0
    merge_results = []

    fileno_map = fetch_fileno_map(pdf_dict.keys())

    for unique_number, documents in pdf_dict.items():
        merged, output_path = merge_documents(unique_number, documents, output_folder)
        if merged:
            merge_results.append((unique_number, True))
            fileno = fileno_map.get(unique_number)
            if fileno:
                new_name = os.path.join(output_folder, f"{fileno}-doc seq.pdf")
                safe_rename(output_path, new_name)
        else:
            merge_results.append((unique_number, False))

        processed += len(documents)
        progress_callback(progress_queue, min(processed / total_documents, 1.0))

    write_results_excel(merge_results, output_folder)

    try:
        os.startfile(output_folder)
    except Exception:
        pass


# ----------------- ENTRY -----------------

if __name__ == "__main__":
    input_folder = r"P:\Users\Steven Cox\Projects\test folders\Rsg Doc merger\In"
    output_folder = r"P:\Users\Steven Cox\Projects\test folders\Rsg Doc merger\Out"
    user_numbers = ["123456789", "987654321"]  # Example 9-digit numbers
    main(input_folder, output_folder, user_numbers)
